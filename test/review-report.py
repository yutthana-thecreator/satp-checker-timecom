# สร้างรายงาน Excel จากผลตัดสินรูปอัตโนมัติของทุกไซต์ (test/out/review_all.json ที่เบราว์เซอร์ POST มา)
# ใช้: python test/review-report.py  → test/out/photo_review_<date>.xlsx + สรุปบน stdout
import json, sys, io, os, collections, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

here = os.path.dirname(os.path.abspath(__file__))
d = json.load(open(os.path.join(here, "out", "review_all.json"), encoding="utf-8"))
rows, sites = d["rows"], d["sites"]
topic_of = lambda r: (r["section"] + " " + r["topic"]) if r["topic"] else r["section"]

wb = Workbook()
head_fill = PatternFill("solid", fgColor="DDE6F2"); bold = Font(bold=True)
def sheet(ws, header, data, widths=None):
    ws.append(header)
    for c in ws[1]: c.font = bold; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in data: ws.append(r)
    for i, w in enumerate(widths or [], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# 1) ต่อไซต์
ws = wb.active; ws.title = "ต่อไซต์"
site_rows = []
for s in sites:
    if s.get("error"): site_rows.append([s["site"], s["profile"], "", "", "", "", s["error"]]); continue
    site_rows.append([s["site"], s["profile"], s["total"], s["auto"], s["need"], round(s["auto"] / s["total"] * 100) if s["total"] else 0, "; ".join(s.get("topicNotes", []))])
sheet(ws, ["ไซต์", "โปรไฟล์", "รูปทั้งหมด", "ผ่านอัตโนมัติ", "ต้องให้ ROM ดู", "% ผ่านอัตโนมัติ", "จำนวนรูปต่างจากตัวอย่าง"], site_rows, [10, 10, 12, 14, 14, 14, 80])

# 2) ต่อโปรไฟล์ × หัวข้อ
agg = collections.defaultdict(lambda: {"n": 0, "need": 0, "sims": []})
for r in rows:
    a = agg[(r["profile"], topic_of(r))]; a["n"] += 1
    if r["status"] == "need": a["need"] += 1
    if r["sim"] is not None: a["sims"].append(r["sim"])
tp = []
for (p, t), a in sorted(agg.items()):
    sims = sorted(a["sims"]); med = sims[len(sims) // 2] if sims else None
    tp.append([p, t, a["n"], a["n"] - a["need"], a["need"], round((a["n"] - a["need"]) / a["n"] * 100), round(med * 100) if med is not None else ""])
ws = wb.create_sheet("ต่อหัวข้อ")
sheet(ws, ["โปรไฟล์", "หัวข้อ", "รูป", "ผ่านอัตโนมัติ", "ต้องให้ ROM ดู", "% ผ่าน", "ความคล้ายกลาง %"], tp, [10, 60, 8, 14, 14, 8, 14])

# 3) รูปที่ต้องให้ ROM ดู — จัดกลุ่มตามสาเหตุ
def reason(note):
    n = note or ""
    if "ไซต์ตัวอย่าง" in n and "รูปเดียวกับ" in n: return "นำรูปไซต์อื่นมาใช้ (ไฟล์เดียวกัน)"
    if "รูปเดียวกับ" in n: return "รูปซ้ำภายในชุดที่โหลด"
    if "นำมาใช้ซ้ำ" in n or "เหมือนรูป" in n: return "รูปซ้ำ/นำรูปไซต์อื่นมาใช้"
    if "เบลอ" in n: return "ภาพเบลอ"
    if "น่าจะเป็นรูปของ" in n: return "วางผิด section"
    if "ยังไม่มีรูปอ้างอิง" in n: return "ไม่มีรูปอ้างอิงของหัวข้อ"
    if "เคยถูก Reject" in n: return "คล้ายรูปที่เคย Reject"
    if "คล้ายรูปอ้างอิงเพียง" in n: return "คล้ายไม่ถึงเกณฑ์"
    return "อื่นๆ"
need = [r for r in rows if r["status"] == "need"]
ws = wb.create_sheet("ต้องให้ ROM ดู")
sheet(ws, ["ไซต์", "โปรไฟล์", "หัวข้อ", "หน้า", "รูปที่", "ความคล้าย %", "สาเหตุ", "รายละเอียดจากระบบ", "การตัดสินใจ ROM (ถอดออก / เก็บไว้)"],
      [[r["site"], r["profile"], topic_of(r), r["page"], r["idx"] + 1, round(r["sim"] * 100) if r["sim"] is not None else "", reason(r["note"]), r["note"], ""] for r in sorted(need, key=lambda r: (reason(r["note"]), r["site"], r["page"]))],
      [10, 10, 55, 6, 6, 12, 26, 90, 30])

# 4) ทุกรูป
ws = wb.create_sheet("ทุกรูป")
sheet(ws, ["ไซต์", "โปรไฟล์", "หัวข้อ", "หน้า", "รูปที่", "ผล", "ความคล้าย %", "รายละเอียด"],
      [[r["site"], r["profile"], topic_of(r), r["page"], r["idx"] + 1, {"auto": "ผ่านอัตโนมัติ", "need": "ต้องให้ ROM ดู", "decided": "ROM ตัดสินแล้ว"}[r["status"]], round(r["sim"] * 100) if r["sim"] is not None else "", r["note"]] for r in rows],
      [10, 10, 55, 6, 6, 16, 12, 90])

out = os.path.join(here, "out", f"photo_review_{datetime.date.today().isoformat()}.xlsx")
wb.save(out)

# สรุป stdout
total = len(rows); n_need = len(need); n_auto = sum(1 for r in rows if r["status"] == "auto")
print(f"FILE {out}")
print(f"sites={len(sites)} photos={total} auto={n_auto} ({n_auto/total*100:.1f}%) need={n_need} ({n_need/total*100:.1f}%)")
print("by reason:");
for k, v in collections.Counter(reason(r["note"]) for r in need).most_common(): print(f"  {k}: {v}")
print("by profile:")
for p in sorted({r['profile'] for r in rows}):
    pr = [r for r in rows if r["profile"] == p]; pn = sum(1 for r in pr if r["status"] == "need")
    print(f"  {p}: photos={len(pr)} need={pn} ({pn/len(pr)*100:.0f}%)")
print("worst topics (need% , n>=15):")
for (p, t), a in sorted(agg.items(), key=lambda kv: -kv[1]["need"] / kv[1]["n"]):
    if a["n"] >= 15 and a["need"]: print(f"  {p} {t}: need {a['need']}/{a['n']} ({a['need']/a['n']*100:.0f}%)")
print("sites with most need:")
for s in sorted([s for s in sites if not s.get("error")], key=lambda s: -s["need"])[:10]: print(f"  {s['site']} ({s['profile']}): need {s['need']}/{s['total']}")
